import io
import html
import os
from datetime import date

import requests
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

API_URL = os.getenv("LEGALEASE_API_URL", "http://127.0.0.1:8000/generate")


def fallback_document(doc_type, parties, terms, effective_date):
    clauses = [x.strip() for x in terms.split(";") if x.strip()]
    items = "\n".join(f"{i}. {x}" for i, x in enumerate(clauses, 1)) or "1. The parties will agree on applicable terms in writing."
    return f"""{doc_type.upper()}\n\nThis {doc_type} is effective from {effective_date}.\n\nPARTIES\n{parties}\n\nTERMS AND CONDITIONS\n{items}\n\nCONFIDENTIALITY\nEach party must protect confidential information disclosed for this arrangement.\n\nSIGNATURES\n\n______________________________        ______________________________\nAuthorized representative                Authorized representative\n\nImportant: This draft is informational and should be reviewed by a qualified legal professional before use."""


def create_docx(text, title):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(11)
    heading = doc.add_heading(title.upper(), 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for block in text.split("\n\n"):
        if block.isupper() and len(block) < 60:
            doc.add_heading(block.title(), level=1)
        else:
            doc.add_paragraph(block)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "LegalEase | AI-assisted legal document draft"
    out = io.BytesIO(); doc.save(out); return out.getvalue()


class LegalPDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 13); self.set_text_color(25, 55, 109)
        self.cell(0, 9, "LegalEase", align="C", new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(25, 55, 109); self.line(10, 21, 200, 21); self.ln(5)
    def footer(self):
        self.set_y(-16); self.set_font("Helvetica", "I", 8); self.set_text_color(100, 100, 100)
        self.cell(0, 8, f"LegalEase | Page {self.page_no()} | AI-assisted legal document draft", align="C")


def create_pdf(text, title):
    pdf = LegalPDF(); pdf.set_auto_page_break(True, 20); pdf.add_page()
    pdf.set_font("Helvetica", "B", 16); pdf.set_text_color(20, 30, 55); pdf.multi_cell(0, 9, title.upper(), align="C"); pdf.ln(4)
    for line in text.splitlines():
        if line.strip().isupper() and len(line.strip()) < 60:
            pdf.ln(3); pdf.set_font("Helvetica", "B", 11); pdf.set_text_color(25, 55, 109)
        else:
            pdf.set_font("Helvetica", "", 10); pdf.set_text_color(30, 30, 30)
        pdf.multi_cell(0, 6, line or " ")
    return bytes(pdf.output())


st.set_page_config(page_title="LegalEase", page_icon="⚖️", layout="wide")
st.markdown("""<style>.stApp{background:#f5f7fb}.hero{padding:1.4rem 0 .6rem}.badge{color:#355cc9;font-weight:700;letter-spacing:.08em}.notice{background:#fff7e6;border-left:4px solid #e5a21a;padding:12px;border-radius:6px}</style>""", unsafe_allow_html=True)
st.markdown("<div class='hero'><div class='badge'>AI-POWERED LEGAL DRAFTING</div><h1>⚖️ LegalEase</h1><p>Create clear, editable legal-document drafts in minutes.</p></div>", unsafe_allow_html=True)
st.markdown("<div class='notice'>LegalEase provides drafting assistance, not legal advice. Have a qualified legal professional review documents before signing.</div>", unsafe_allow_html=True)

with st.form("document_form"):
    col1, col2 = st.columns(2)
    with col1:
        doc_type = st.selectbox("Document type", ["Non-Disclosure Agreement", "Employment Contract", "Lease Agreement", "Service Agreement", "Freelance Work Contract", "Custom Agreement"])
        parties = st.text_area("Parties involved", placeholder="Jane Doe (Service Provider), TechNova Inc. (Client)", height=120)
    with col2:
        effective_date = st.date_input("Effective date", value=date.today()).strftime("%B %d, %Y")
        terms = st.text_area("Terms & conditions", placeholder="Payment within 30 days; Maintain confidentiality; Either party may terminate with 15 days notice", height=120)
    submitted = st.form_submit_button("Generate document", type="primary", use_container_width=True)

if submitted:
    if not parties.strip(): st.error("Please enter the parties involved.")
    else:
        payload = {"document_type": doc_type, "parties": parties, "terms": terms, "effective_date": effective_date}
        try:
            response = requests.post(API_URL, json=payload, timeout=45); response.raise_for_status()
            st.session_state.document = response.json()["document"]
            st.session_state.generation_note = "Generated through the LegalEase API."
        except requests.RequestException:
            st.session_state.document = fallback_document(**payload)
            st.session_state.generation_note = "API is offline, so LegalEase created a local template draft. Start FastAPI to enable Gemini generation."
        st.session_state.doc_type = doc_type

if "document" in st.session_state:
    st.divider(); st.subheader("Review and edit")
    st.caption(st.session_state.get("generation_note", ""))
    edited = st.text_area("Document draft", st.session_state.document, height=430)
    st.session_state.document = edited
    st.markdown("### Preview")
    st.markdown(f"<div style='background:white;padding:28px;border-radius:12px;border:1px solid #e5e7eb;white-space:pre-wrap;font-family:Georgia,serif;line-height:1.65'>{html.escape(edited)}</div>", unsafe_allow_html=True)
    a, b, c = st.columns(3)
    safe_name = "legalease_" + st.session_state.doc_type.lower().replace(" ", "_")
    a.download_button("Download TXT", edited.encode("utf-8"), f"{safe_name}.txt", "text/plain", use_container_width=True)
    b.download_button("Download DOCX", create_docx(edited, st.session_state.doc_type), f"{safe_name}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
    c.download_button("Download PDF", create_pdf(edited, st.session_state.doc_type), f"{safe_name}.pdf", "application/pdf", use_container_width=True)
